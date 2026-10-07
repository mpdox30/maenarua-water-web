# -*- coding: utf-8 -*-
"""สคริปต์สลับ production -> สูตร spill ใหม่ + โมเดล candidate  (DEFAULT = DRY-RUN, ไม่เขียนอะไร)

ใช้:  python 05_cutover.py                 # แสดงแผน + ตรวจ precondition
      python 05_cutover.py --apply         # ทำจริง (สำรองไฟล์ .bak_before_newspill_<ts> ก่อนแก้ทุกไฟล์)
      python 05_cutover.py --apply --patch-days 14 --ledger-dir "D:\\WMB_Phayao\\01_raw_data\\Reservoirs\\บัญชีน้ำ"
รันบน Windows (path จริง) หลังรัน recompute_corrected_ledger.py แล้ว ledger เป็น offset 0 ล่าสุด

ขั้นตอน
 1 reservoir_water_balance.py : compute_spillway_overflow_m3 -> critical-flow+Manning (n=0.017) ตามตาราง friction curve, offset 0,
                                time-weighted 1 ชม./ค่า (เหมือน ledger); คัดลอก curve CSV ไป Reservoirs/reference/
 2 official xlsx ย้อนหลัง N วัน : เขียนทับคอลัมน์ D (Inflow) และ F (Spill) ด้วยค่าจาก ledger (ข้ามวัน fallback) -- ใช้ flatten ของ writer
 3 โมเดล                        : candidate_regressors_newspill.pkl -> active/deployment_regressors_no_stage1.pkl
 4 data_pipeline.py             : RAIN_FORECAST_BINARY_HORIZONS={6:3.1}, RESERVOIR_INFLOW_BIAS_CORRECTION_M3={}
 5 Training CSV                 : คัดลอก Training_Values_newspill_offset0 -> active/Training_Values_Nofct_7day_Extended_lagfeat.csv (เดิมสำรองไว้)
 6 model_metadata / deployment_model_choice : เพิ่ม note + ตั้ง bias=0, rain flag ตาม candidate
หลัง apply: รัน orchestration 1 ครั้ง, ตรวจ forecast, git commit; ล้าง/ปรับ forecast_accuracy_log basis เองถ้าต้องการ (ไม่ทำในสคริปต์นี้)
"""
import argparse, json, re, shutil, datetime as dt, sys
from pathlib import Path
HERE=Path(__file__).resolve().parent
REPO=Path(__file__).resolve().parents[3]            # .../maenaruea-water-web
PIPE=REPO/"01_data/scripts and code/pipeline"
ACTIVE=REPO/"01_data/scripts and code/Reservoir_inflow/active"
REFDIR=REPO/"01_data/Reservoirs/reference"
TS=dt.datetime.now().strftime("%Y%m%d_%H%M%S")
ap=argparse.ArgumentParser();ap.add_argument("--apply",action="store_true");ap.add_argument("--patch-days",type=int,default=14)
ap.add_argument("--ledger-dir",default=r"D:\WMB_Phayao\01_raw_data\Reservoirs\บัญชีน้ำ");a=ap.parse_args()
def bak(p):
    b=p.with_name(p.name+f".bak_before_newspill_{TS}");print("  backup",b.name)
    if a.apply: shutil.copy2(p,b)
def say(s): print(("[APPLY] " if a.apply else "[DRY]   ")+s)

NEW_FN='''def _load_spillway_friction_curve():
    """ตาราง critical-flow + Manning friction (n=0.017): head (m) -> q ต่อความกว้าง 1 m (m2/s). เพิ่ม 2026-10 ตามนโยบายสูตรใหม่"""
    global _spill_curve_cache
    if _spill_curve_cache is None:
        heads, qs = [], []
        with open(SPILLWAY_FRICTION_CURVE_CSV, "r", encoding="utf-8") as f:
            next(f)
            for line in f:
                h, q = line.strip().split(",")
                heads.append(float(h)); qs.append(float(q))
        _spill_curve_cache = (heads, qs)
    return _spill_curve_cache


def compute_spillway_overflow_m3(hourly_levels_msl: list[float]) -> float:
    """
    ปริมาณน้ำล้นสปิลเวย์รายวัน (m3) จากระดับน้ำรายชั่วโมง -- สูตรใหม่ (2026-10): critical-flow + Manning friction
    (n=0.017) interpolate จาก spillway_friction_curve_n017.csv, offset ระดับน้ำ = 0 (ไม่หัก 0.155), L = 30 m,
    crest 489.545 ; แต่ละค่าแทน 1 ชั่วโมง (rectangle) -- สอดคล้องกับ recompute_corrected_ledger.py
    (สูตรเดิม weir C*L*H^1.5 + offset 0.155 ถูกแทนที่ที่นี่)
    """
    import numpy as _np
    heads, qs = _load_spillway_friction_curve()
    head = _np.clip(_np.asarray(hourly_levels_msl, dtype=float) - SPILLWAY_CREST_LEVEL_MSL, 0.0, None)
    q_m3_s = _np.interp(head, heads, qs) * SPILLWAY_WEIR_LENGTH_M
    return float((q_m3_s * 3600.0).sum())
'''
# ---- 0 precondition
curve_src=HERE/"spillway_friction_curve_n017.csv"
print("== precondition");
for p in (curve_src,HERE/"candidate_regressors_newspill.pkl",HERE/"candidate_metadata_newspill.json",HERE/"Training_Values_newspill_offset0_20261007.csv",PIPE/"reservoir_water_balance.py",PIPE/"data_pipeline.py"):
    print("  ",("OK  " if p.exists() else "MISSING"),p.name)
# ---- 1
rwb=PIPE/"reservoir_water_balance.py";src=rwb.read_text(encoding="utf-8");print("== 1 reservoir_water_balance.py")
pat=re.compile(r"def compute_spillway_overflow_m3\(.*?\n    return total_m3\n",re.S)
assert pat.search(src),"ไม่พบฟังก์ชันเดิม"
if "_load_spillway_friction_curve" in src: print("  already applied")
else:
    new=pat.sub(lambda m:NEW_FN,src)
    new=new.replace("SPILLWAY_LEVEL_ADJ_OFFSET_M = 0.155","SPILLWAY_LEVEL_ADJ_OFFSET_M = 0.155   # (เลิกใช้ในสูตรใหม่ 2026-10 -- เก็บไว้อ้างอิง)\nSPILLWAY_FRICTION_CURVE_CSV = REFERENCE_DIR / \"spillway_friction_curve_n017.csv\"\n_spill_curve_cache = None",1)
    bak(rwb);say("เขียน compute_spillway_overflow_m3 ใหม่ + คัดลอก curve ไป Reservoirs/reference/")
    if a.apply:
        rwb.write_text(new,encoding="utf-8");REFDIR.mkdir(parents=True,exist_ok=True);shutil.copy2(curve_src,REFDIR/curve_src.name)
# ---- 2
print("== 2 official xlsx ย้อนหลัง",a.patch_days,"วัน")
import openpyxl
sys.path.insert(0,str(PIPE))
led={}
for f in Path(a.ledger_dir).glob("*_MNR.xlsx"):
    y,mn,_=f.name.split("_");mnum=dt.datetime.strptime(mn,"%B").month
    ws=openpyxl.load_workbook(f,data_only=True).active
    for r in ws.iter_rows(min_row=6):
        if isinstance(r[0].value,int) and r[3].value is not None and not str(r[11].value or "").startswith("ใช้ค่าเดิม"):
            led[dt.date(int(y),mnum,r[0].value)]=(r[3].value,r[5].value)
last=max(led) if led else None
if last:
    days=[last-dt.timedelta(days=i) for i in range(a.patch_days)][::-1]
    import reservoir_official_file_writer as w
    by={}
    for d in days:
        if d in led: by.setdefault(w.official_file_path(d),[]).append(d)
    for p,ds in by.items():
        wb0=openpyxl.load_workbook(p,data_only=True).active
        for d in ds:
            r=d.day+5;print("  ",d,"D",round(wb0.cell(r,4).value or 0),"->",round(led[d][0]),"| F",round(wb0.cell(r,6).value or 0),"->",round(led[d][1]))
        if a.apply:
            bak(p);wb=w._flatten_formulas_to_values(p);ws=wb["บัญชีน้ำ"] if "บัญชีน้ำ" in wb.sheetnames else wb.active
            for d in ds: ws.cell(d.day+5,4).value=led[d][0];ws.cell(d.day+5,6).value=led[d][1]
            wb.save(p)
# ---- 3-6
print("== 3 โมเดล / 5 training CSV")
for s,d in ((HERE/"candidate_regressors_newspill.pkl",ACTIVE/"deployment_regressors_no_stage1.pkl"),(HERE/"Training_Values_newspill_offset0_20261007.csv",ACTIVE/"Training_Values_Nofct_7day_Extended_lagfeat.csv")):
    bak(d);say(f"{s.name} -> {d.name}")
    if a.apply: shutil.copy2(s,d)
print("== 4 data_pipeline.py");dp=PIPE/"data_pipeline.py";t=dp.read_text(encoding="utf-8")
t2=t.replace("RAIN_FORECAST_BINARY_HORIZONS = {3: 27.3, 6: 3.1, 7: 61.0}","RAIN_FORECAST_BINARY_HORIZONS = {6: 3.1}   # 2026-10 retrain สูตร spill ใหม่ -- ใช้เฉพาะ H6")
t2=t2.replace("RESERVOIR_INFLOW_BIAS_CORRECTION_M3 = {3: 610.9, 4: 4241.9, 5: 2605.5, 6: 6583.5, 7: 13153.2}","RESERVOIR_INFLOW_BIAS_CORRECTION_M3 = {}   # 2026-10 retrain สูตร spill ใหม่ -- ยกเลิก (ไม่ช่วยใน walk-forward); refit จาก live log ภายหลัง")
print("  changed" if t2!=t else "  already applied / no match")
if a.apply and t2!=t: bak(dp);dp.write_text(t2,encoding="utf-8")
print("== 6 metadata")
mm=ACTIVE/"model_metadata.json";dc=ACTIVE/"deployment_model_choice.json"
if a.apply:
    for p in (mm,dc): bak(p)
    m=json.loads(mm.read_text(encoding="utf-8"));m["retrain_newspill_2026-10"]=dict(dataset="Training_Values_newspill_offset0_20261007.csv",candidate_dir="01_data/experiments/daily_newspill_retrain_20261007",rain_forecast_horizons=[6],bias_correction="none")
    mm.write_text(json.dumps(m,indent=2,ensure_ascii=False),encoding="utf-8")
    c=json.loads(dc.read_text(encoding="utf-8"))
    for h,v in c.items():
        if isinstance(v,dict): v["updated_2026-10_newspill"]=dict(bias_correction_m3_per_day=0.0,uses_rain_forecast_binary=(h=="6"),rain_forecast_threshold_mm=(3.1 if h=="6" else None),training_data_source="Training_Values_newspill_offset0_20261007.csv",note="test_nse/walkforward เดิมเป็นของรุ่นก่อน -- ดู daily_newspill_retrain_20261007/cv_*summary.csv")
    dc.write_text(json.dumps(c,indent=2,ensure_ascii=False),encoding="utf-8")
print("\nDRY-RUN เสร็จ -- ใส่ --apply เพื่อทำจริง" if not a.apply else "\nAPPLY เสร็จ -- รัน orchestration/ทดสอบ แล้ว commit")
